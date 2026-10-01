# -*- coding: utf-8 -*-
"""Проверка готового отчёта (.docx) на соответствие требованиям нормоконтроля."""
import os
import re
import sys

from docx import Document
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
DOCX = os.path.join(ROOT, "Отчёт", "Отчёт_Практика_АБД_Поликлиника_Сагадиев.docx")

problems = []


def check(cond, message):
    if not cond:
        problems.append(message)


def cm(value):
    return round(value.cm, 2)


def near(a, b, tol=0.05):
    return abs(a - b) <= tol


def main():
    doc = Document(DOCX)
    secs = doc.sections
    check(len(secs) == 3, "секций документа: %d (ожидалось 3)" % len(secs))

    if len(secs) == 3:
        title, toc, body = secs
        expected = [
            ("титульный лист", title, (2.0, 2.0, 2.5, 1.5), False),
            ("СОДЕРЖАНИЕ", toc, (1.5, 5.5, 3.0, 1.5), True),
            ("основной текст", body, (1.5, 2.5, 2.5, 1.5), True),
        ]
        for name, sec, (top, bottom, left, right), with_frame in expected:
            check(near(cm(sec.top_margin), top), "%s: верхнее поле %.2f см (нужно %.1f)"
                  % (name, cm(sec.top_margin), top))
            check(near(cm(sec.bottom_margin), bottom), "%s: нижнее поле %.2f см (нужно %.1f)"
                  % (name, cm(sec.bottom_margin), bottom))
            check(near(cm(sec.left_margin), left), "%s: левое поле %.2f см (нужно %.1f)"
                  % (name, cm(sec.left_margin), left))
            check(near(cm(sec.right_margin), right), "%s: правое поле %.2f см (нужно %.1f)"
                  % (name, cm(sec.right_margin), right))
            xml = sec._sectPr.xml
            has_borders = "pgBorders" in xml
            check(has_borders == with_frame, "%s: рамка ЕСКД %s"
                  % (name, "отсутствует" if with_frame else "не должна быть"))
            if with_frame:
                check('w:offsetFrom="page"' in xml,
                      "%s: рамка должна быть привязана к странице" % name)

    # титульный лист: реквизиты (текст абзацев и таблиц первой страницы)
    title_parts = [p.text for p in doc.paragraphs[:60]]
    for table in doc.tables[:2]:
        for row in table.rows[:6]:
            for cell in row.cells:
                title_parts.append(cell.text)
    title_text = "\n".join(title_parts).replace("\u00a0", " ")
    for needle in ("Уфимский государственный колледж технологии и дизайна",
                   "Р-25-29", "Сагадиев", "Габитова", "09.02.07", "2026"):
        check(needle in title_text, "титульный лист: не найдено «%s»" % needle)

    # таблицы: подпись «Таблица N – …» и строка нумерации столбцов
    table_captions = [p.text.strip() for p in doc.paragraphs
                      if re.match(r"^Таблица\s+\d+\s+–\s+", p.text.strip())]
    content_tables = doc.tables[1:]          # первая таблица — штамп титульного листа
    check(len(content_tables) == 17,
          "таблиц в отчёте: %d (ожидалось 17 содержательных)" % len(content_tables))
    check(len(table_captions) == 17, "подписей «Таблица N – …»: %d" % len(table_captions))

    numbered = 0
    for table in content_tables:
        if len(table.rows) < 2:
            continue
        row = [c.text.strip() for c in table.rows[1].cells]
        if row and row[0] == "1" and all(row[i] == str(i + 1) for i in range(len(row))):
            numbered += 1
    check(numbered == 17, "таблиц со строкой нумерации столбцов 1..N: %d" % numbered)

    # рисунки
    fig_captions = [p.text.strip() for p in doc.paragraphs
                    if re.match(r"^Рисунок\s+\d+\s+–\s+", p.text.strip())]
    check(len(fig_captions) == 10, "подписей «Рисунок N – …»: %d" % len(fig_captions))

    # формулы (нумерация вида «(4.1)» в правом крае)
    formula_numbers = [p.text.strip() for p in doc.paragraphs
                       if re.search(r"\(\d+\.\d+\)\s*$", p.text.strip())]
    check(len(formula_numbers) >= 3, "пронумерованных формул: %d" % len(formula_numbers))

    # основной текст: параметры абзацев
    body_paras = []
    for p in doc.paragraphs:
        pf = p.paragraph_format
        if not p.text.strip() or pf.first_line_indent is None:
            continue
        run = next((r for r in p.runs if r.text.strip()), None)
        if run is None or run.font.size is None:
            continue
        if abs(pf.first_line_indent.cm - 1.25) < 0.01 and run.font.size.pt == 14:
            body_paras.append(p)
    check(len(body_paras) > 50, "абзацев основного текста (14 пт, отступ 1,25 см): %d"
          % len(body_paras))
    for p in body_paras:
        pf = p.paragraph_format
        run = next(r for r in p.runs if r.text.strip())
        check(run.font.name in ("Times New Roman", None),
              "шрифт абзаца «%s…»: %s" % (p.text[:25], run.font.name))
        check(near(pf.line_spacing or 0, 1.5),
              "межстрочный интервал абзаца «%s…»: %s" % (p.text[:25], pf.line_spacing))

    # отсутствие недопустимых имён в квадратных скобках в листингах SQL
    for p in doc.paragraphs:
        for m in re.finditer(r"\[([^\]\n]+)\]", p.text):
            name = m.group(1)
            if any(ch in name for ch in ".,!`[]"):
                problems.append("листинг SQL: недопустимое имя [%s]" % name)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for m in re.finditer(r"\[([^\]\n]+)\]", cell.text):
                    name = m.group(1)
                    if any(ch in name for ch in ".,!`[]"):
                        problems.append("таблица отчёта: недопустимое имя [%s]" % name)

    # шрифт основного текста в стиле
    style = doc.styles["Body Text"] if "Body Text" in [s.name for s in doc.styles] else None
    if style is not None:
        check(style.font.size is None or style.font.size.pt == 14,
              "размер шрифта стиля основного текста: %s" % style.font.size)

    if problems:
        print("ОБНАРУЖЕНЫ ПРОБЛЕМЫ (%d):" % len(problems))
        for p in problems[:40]:
            print("  -", p)
        return 1
    print("Отчёт проверен: 3 секции, 17 таблиц, 10 рисунков, формулы, поля и рамки — в норме.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
