# -*- coding: utf-8 -*-
"""Генератор книги MS Excel Диаграмма_ЗагрузкаВрачей.xlsx (Приложение В)."""
import os
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import BarChart, Reference

from db_schema import build_data

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))


def main():
    data = build_data()
    wb = openpyxl.Workbook()

    font_title = Font(name="Arial", size=12, bold=True)
    font_hdr = Font(name="Arial", size=10, bold=True)
    font_cell = Font(name="Arial", size=10, bold=False)
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_hdr = PatternFill("solid", fgColor="D9D9D9")

    # --- Лист 1: Загрузка врачей + диаграмма ---
    ws = wb.active
    ws.title = "Загрузка врачей"
    ws["A1"] = "Загрузка врачей поликлиники по специальностям (сентябрь 2026 г.)"
    ws["A1"].font = font_title

    headers = ["Специальность", "Всего талонов", "Принято", "Запланировано", "Неявка", "Отменено", "Загрузка, %"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = font_hdr
        c.fill = fill_hdr
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    spec_names = {s["КодСпециальности"]: s["НазваниеСпециальности"] for s in data["Специальности"]}
    doc_spec = {d["КодВрача"]: d["КодСпециальности"] for d in data["Врачи"]}
    st = defaultdict(lambda: {"Всего": 0, "Проведён": 0, "Запланирован": 0, "Неявка": 0, "Отменён": 0})
    for z in data["ЗаписиНаПриём"]:
        name = spec_names[doc_spec[z["КодВрача"]]]
        st[name]["Всего"] += 1
        st[name][z["СтатусЗаписи"]] += 1

    order = [s["НазваниеСпециальности"] for s in sorted(data["Специальности"], key=lambda x: x["ПорядокОтображения"])]
    for idx, name in enumerate(order, start=4):
        s = st[name]
        vsego = s["Всего"]
        prin = s["Проведён"]
        pct = round(100.0 * prin / vsego, 1) if vsego else 0.0
        row_vals = [name, vsego, prin, s["Запланирован"], s["Неявка"], s["Отменён"], pct]
        for col, val in enumerate(row_vals, start=1):
            c = ws.cell(row=idx, column=col, value=val)
            c.font = font_cell
            c.border = border
            c.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")

    tot_row = 4 + len(order)
    ws.cell(row=tot_row, column=1, value="Итого").font = font_hdr
    ws.cell(row=tot_row, column=1).border = border
    for col in range(2, 7):
        col_letter = openpyxl.utils.get_column_letter(col)
        c = ws.cell(row=tot_row, column=col, value=f"=SUM({col_letter}4:{col_letter}{tot_row-1})")
        c.font = font_hdr
        c.border = border
        c.alignment = Alignment(horizontal="center")
    c_pct = ws.cell(row=tot_row, column=7, value=f"=ROUND(C{tot_row}/B{tot_row}*100,1)")
    c_pct.font = font_hdr
    c_pct.border = border
    c_pct.alignment = Alignment(horizontal="center")

    widths = [28, 16, 14, 16, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Нативная диаграмма Excel по п. 1.7 нормоконтроля
    chart = BarChart()
    chart.type = "col"
    chart.style = 2
    chart.title = "Загрузка врачей поликлиники по специальностям (сентябрь 2026 г.)"
    chart.y_axis.title = "Количество талонов"
    chart.x_axis.title = "Специальность"
    chart.legend.position = "b"
    chart.width = 18
    chart.height = 11

    data_ref = Reference(ws, min_col=3, min_row=3, max_col=4, max_row=tot_row - 1)
    cats_ref = Reference(ws, min_col=1, min_row=4, max_row=tot_row - 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "A15")

    # --- Лист 2: Журнал записей (выгрузка Q03_ЖурналЗаписей) ---
    ws2 = wb.create_sheet("Журнал записей")
    hdr2 = ["№ талона", "Дата приёма", "Время", "Пациент", "Врач", "Специальность", "Статус"]
    for col, h in enumerate(hdr2, start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = font_hdr
        c.fill = fill_hdr
        c.border = border
    pats = {p["КодПациента"]: f"{p['Фамилия']} {p['Имя'][0]}. {p['Отчество'][0]}." for p in data["Пациенты"]}
    docs = {d["КодВрача"]: f"{d['Фамилия']} {d['Имя'][0]}. {d['Отчество'][0]}." for d in data["Врачи"]}
    for r_i, z in enumerate(sorted(data["ЗаписиНаПриём"], key=lambda x: (x["ДатаПриёма"], x["ВремяПриёма"])), start=2):
        vals = [
            z["НомерТалона"],
            z["ДатаПриёма"].strftime("%d.%m.%Y"),
            z["ВремяПриёма"].strftime("%H:%M"),
            pats[z["КодПациента"]],
            docs[z["КодВрача"]],
            spec_names[doc_spec[z["КодВрача"]]],
            z["СтатусЗаписи"],
        ]
        for col, v in enumerate(vals, start=1):
            c = ws2.cell(row=r_i, column=col, value=v)
            c.font = font_cell
            c.border = border
    for i, w in enumerate([18, 14, 10, 24, 24, 26, 16], start=1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    out_path = os.path.join(ROOT, "Диаграмма_ЗагрузкаВрачей.xlsx")
    wb.save(out_path)
    print("  книга Excel сохранена:", out_path)


if __name__ == "__main__":
    main()
